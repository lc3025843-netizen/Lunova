"""DOCX extraction and export helpers."""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO

from docx import Document


def extract_docx_text(file_bytes: bytes) -> tuple[str, int]:
    doc = Document(BytesIO(file_bytes))
    chunks: list[str] = []
    count = 0

    for paragraph in doc.paragraphs:
        value = paragraph.text.strip()
        if value:
            chunks.append(value)
            count += 1

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    value = paragraph.text.strip()
                    if value:
                        chunks.append(value)
                        count += 1

    return "\n\n".join(chunks), count


def export_revised_docx(original_bytes: bytes, revised_text: str) -> bytes:
    """
    Create a safe revised copy.

    V1 intentionally writes the revised body into a new section after a divider instead of
    destructively replacing runs in the original. This preserves the user's original formatting,
    tables, references, images, headers, and footers while providing an editable revised version.
    """
    source = Document(BytesIO(original_bytes))
    source.add_page_break()
    heading = source.add_paragraph()
    run = heading.add_run("Versión revisada por Lunova")
    run.bold = True

    for block in [p.strip() for p in revised_text.split("\n\n") if p.strip()]:
        source.add_paragraph(block)

    output = BytesIO()
    source.save(output)
    return output.getvalue()
